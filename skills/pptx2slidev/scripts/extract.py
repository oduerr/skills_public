"""Extract the visible slides of a .pptx into a text dump, a JSON index and a media folder.

Usage: python extract.py deck.pptx OUTDIR

Writes
  OUTDIR/dump.md      one block per visible slide: shapes with position (% of slide),
                      text, tables, images, notes, links, animations
  OUTDIR/slides.json  per visible slide: number, original number, title, all text
                      (full XML walk, used by compare.py for the lost-words check)
  OUTDIR/media/       images as sNN_<shape-id>.<ext> (NN = visible slide number)

  OUTDIR/pictures.json per picture: slide, file, raw file if cropped, crop, rotation, stretch
                      (used by check_images.py)

Pictures are saved as PowerPoint shows them: a srcRect crop is applied (the uncropped
original is kept as *.raw.<ext> and must not be used: crops often hide private browser
tabs, answers or extra data). Rotation, flips and stretched frames are reported.
Text-less shapes lying over a picture are reported as COVER (often a box that hides
an answer until a click).
Hidden slides are skipped, because PDF exports skip them too.
Text is also collected by walking the full slide XML (a:t and m:t, without
mc:Fallback), because python-pptx misses text inside mc:AlternateContent and
OMML math. If the "raw text" line of a slide has words that the shape list
does not show, look at the reference page: that text must not get lost.
Needs: python-pptx.
"""
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"

src, out = sys.argv[1], sys.argv[2]
os.makedirs(f"{out}/media", exist_ok=True)
prs = Presentation(src)
W, H = prs.slide_width, prs.slide_height
lines = []


def pos(sh):
    try:
        return f"x={100*sh.left/W:.0f}% y={100*sh.top/H:.0f}% w={100*sh.width/W:.0f}% h={100*sh.height/H:.0f}%"
    except Exception:
        return "pos=?"


def omml_text(el):
    return "".join(t.text or "" for t in el.iter(f"{M}t"))


def raw_text(el):
    """All a:t and m:t text below el, skipping mc:Fallback (duplicate of the math)."""
    parts = []

    def walk(e):
        if e.tag == f"{MC}Fallback":
            return
        if e.tag in (f"{A}t", f"{M}t") and e.text:
            parts.append(e.text)
        for c in e:
            walk(c)

    walk(el)
    return " ".join(parts)


def para_text(p):
    parts = []
    for child in p._p:
        tag = child.tag
        if tag == f"{A}r":
            t = "".join(x.text or "" for x in child.iter(f"{A}t"))
            rpr = child.find(f"{A}rPr")
            if rpr is not None and rpr.get("b") == "1" and t.strip():
                t = f"**{t}**"
            parts.append(t)
        elif tag == f"{A}br":
            parts.append(" / ")
        elif tag == f"{A}fld":
            parts.append("".join(x.text or "" for x in child.iter(f"{A}t")))
        else:
            m = list(child.iter(f"{M}oMath"))
            if m:
                parts.append(" [MATH: " + " ".join(omml_text(x) for x in m) + "] ")
    return "".join(parts)


def dump_shape(sh, si, depth=0):
    ind = "  " * depth
    kind = sh.shape_type
    if kind == MSO_SHAPE_TYPE.GROUP:
        lines.append(f"{ind}- GROUP {pos(sh)}")
        for s in sh.shapes:
            dump_shape(s, si, depth + 1)
        return
    if kind == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
        try:
            img = sh.image
            name = re.sub(r"[^A-Za-z0-9_.-]", "_", f"s{si:02d}_{sh.shape_id}.{img.ext}")
            info = save_picture(sh, img, name)
            descr = sh._element.xpath(".//p:cNvPr/@descr")
            alt = f" descr='{descr[0]}'" if descr else ""
            lines.append(f"{ind}- IMAGE media/{info['file']} {pos(sh)} ({img.size[0]}x{img.size[1]}px) name='{sh.name}'{alt}{info['note']}")
            pictures.append({"slide": si, "shape_id": sh.shape_id, "bbox": bbox(sh), **info})
        except Exception as e:
            lines.append(f"{ind}- IMAGE? {sh.name} {pos(sh)} ({e}) -> crop from reference page")
        return
    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        txt = [(p.level, para_text(p)) for p in sh.text_frame.paragraphs]
        if not any(t.strip() for _, t in txt):
            m = list(sh._element.iter(f"{M}oMath"))
            if m:
                lines.append(f"{ind}- MATH {pos(sh)}: " + " | ".join(omml_text(x) for x in m))
            return
        ph = f" placeholder={sh.placeholder_format.type}" if sh.is_placeholder else ""
        lines.append(f"{ind}- TEXT {pos(sh)}{ph}")
        for lvl, t in txt:
            if t.strip():
                lines.append(f"{ind}    {'  '*lvl}* {t}")
        return
    if getattr(sh, "has_table", False) and sh.has_table:
        lines.append(f"{ind}- TABLE {pos(sh)}")
        for r in sh.table.rows:
            lines.append(f"{ind}    | " + " | ".join(c.text.replace(chr(10), ' ') for c in r.cells) + " |")
        return
    if getattr(sh, "has_chart", False) and sh.has_chart:
        lines.append(f"{ind}- CHART {pos(sh)} type={sh.chart.chart_type} -> crop from reference page")
        return
    xml = sh._element.xml
    if "videoFile" in xml or "audioFile" in xml or kind == MSO_SHAPE_TYPE.MEDIA:
        lines.append(f"{ind}- MEDIA/VIDEO {sh.name} {pos(sh)}")
        return
    m = list(sh._element.iter(f"{M}oMath"))
    if m:
        lines.append(f"{ind}- MATH {pos(sh)}: " + " | ".join(omml_text(x) for x in m))
        return
    if kind == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
        lines.append(f"{ind}- OLE {sh.name} {pos(sh)} -> crop from reference page (formula? rewrite in KaTeX)")
        return
    lines.append(f"{ind}- SHAPE {kind} {sh.name} {pos(sh)}")


def bbox(sh):
    try:
        return [sh.left / W, sh.top / H, (sh.left + sh.width) / W, (sh.top + sh.height) / H]
    except Exception:
        return None


def save_picture(sh, img, name):
    """Save the picture as PowerPoint shows it: srcRect crop applied, rotation and
    stretching reported. If cropped, the uncropped original is kept as *.raw.<ext>
    (never use the raw file on a slide: crops often hide private content or answers)."""
    el = sh._element
    note, info = "", {"file": name, "raw": None, "crop": None, "rot": 0, "flip": False, "stretch": None}
    blob = img.blob
    sr = el.find(".//" + A + "srcRect")
    crop = None
    if sr is not None:
        crop = [int(sr.get(k, 0)) / 1000 for k in ("l", "t", "r", "b")]  # percent
        if not any(abs(c) > 0.05 for c in crop):
            crop = None
    xfrm = el.find(".//" + A + "xfrm")
    rot = int(xfrm.get("rot", 0)) / 60000 if xfrm is not None else 0
    flip = xfrm is not None and (xfrm.get("flipH") == "1" or xfrm.get("flipV") == "1")
    try:
        from PIL import Image
        import io
        im = Image.open(io.BytesIO(blob))
        im.load()
        w, h = im.size
        if crop:
            l, t, r, b = crop
            box = (round(w * l / 100), round(h * t / 100), round(w * (1 - r / 100)), round(h * (1 - b / 100)))
            box = (max(0, box[0]), max(0, box[1]), min(w, box[2]), min(h, box[3]))
            raw = name.rsplit(".", 1)[0] + ".raw." + name.rsplit(".", 1)[1]
            with open(f"{out}/media/{raw}", "wb") as f:
                f.write(blob)
            cim = im.crop(box)
            if cim.mode in ("RGBA", "LA", "P") and name.lower().endswith((".jpg", ".jpeg")):
                cim = cim.convert("RGB")
            name = name.rsplit(".", 1)[0] + (".png" if not name.lower().endswith((".jpg", ".jpeg")) else ".jpg")
            cim.save(f"{out}/media/{name}")
            info.update(file=name, raw=raw)
            w, h = cim.size
            note += f" CROPPED in PowerPoint L/T/R/B={l:.0f}/{t:.0f}/{r:.0f}/{b:.0f}% (file is cropped; raw kept as {raw}, do not use it)"
        else:
            with open(f"{out}/media/{name}", "wb") as f:
                f.write(blob)
        frame = sh.width / sh.height if sh.height else 0
        pic = w / h if h else 0
        if frame and pic and abs(frame / pic - 1) > 0.10:
            info["stretch"] = round(frame / pic, 2)
            note += f" STRETCHED: frame ratio {frame:.2f} vs image {pic:.2f} (PowerPoint shows it distorted; reproduce with object-fill or a fixed box)"
    except Exception as e:
        with open(f"{out}/media/{name}", "wb") as f:
            f.write(blob)
        a, b = blob.find(b"%PDF"), blob.rfind(b"%%EOF")
        if a >= 0 and b > a:  # EMF+/WMF with an embedded PDF (e.g. Beamer slides): LibreOffice renders these empty
            pdfname = name.rsplit(".", 1)[0] + ".embedded.pdf"
            with open(f"{out}/media/{pdfname}", "wb") as f:
                f.write(blob[a:b + 5])
            import subprocess
            subprocess.run(["pdftoppm", "-png", "-r", "200", "-singlefile", f"{out}/media/{pdfname}",
                            f"{out}/media/" + name.rsplit(".", 1)[0] + ".embedded"], capture_output=True)
            note += f" EMBEDDED PDF extracted -> media/{name.rsplit('.', 1)[0]}.embedded.png (use this, the reference render may be empty)"
        if crop:
            note += f" CROPPED in PowerPoint L/T/R/B={'/'.join(f'{c:.0f}' for c in crop)}% but could not apply ({e}): crop from the reference page instead"
        info["file"] = name
    if rot:
        note += f" ROTATED {rot:.0f}deg"
    if flip:
        note += " FLIPPED"
    info.update(crop=crop, rot=rot, flip=flip, note=note)
    return info


def covers(slide, si):
    """Text-less shapes that lie over a picture: often boxes that hide an answer until
    a click (exit animation) or blank out something. Reported so the answer is not shown."""
    timing = slide._element.xml
    out_lines = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    pics = [p for p in pictures if p["slide"] == si and p["bbox"]]
    # formula objects (OLE) and charts can be covered too
    for sh in walk(slide.shapes):
        if sh.shape_type in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.CHART) and bbox(sh):
            pics.append({"file": f"(object '{sh.name}')", "bbox": bbox(sh)})

    for sh in walk(slide.shapes):
        if sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.CHART,
                             MSO_SHAPE_TYPE.TABLE) or hasattr(sh, "image") or sh.is_placeholder or sh._element.tag.endswith("graphicFrame"):
            continue
        kind = "COVER"
        sppr = sh._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
        geom = sppr.find(A + "prstGeom") if sppr is not None else None
        if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh._element.tag.endswith("cxnSp") or \
                (geom is not None and re.search(r"line|connector|arc|bracket|brace|arrow", geom.get("prst", ""), re.I)):
            kind = "MARK"  # arrow / line drawn on the picture
        elif sppr is None or sppr.find(A + "noFill") is not None or \
                not any(sppr.find(A + f) is not None for f in ("solidFill", "gradFill", "blipFill", "pattFill")):
            kind = "MARK"  # unfilled outline (circle, frame) highlighting part of the picture
        else:
            fill = sppr.find(A + "solidFill")
            clr = fill[0] if fill is not None and len(fill) else None
            whiteish = clr is not None and (
                (clr.tag == A + "srgbClr" and clr.get("val", "").upper() in ("FFFFFF", "FEFEFE", "FDFDFD", "F2F2F2"))
                or (clr.tag == A + "schemeClr" and clr.get("val") in ("bg1", "lt1", "bg2", "lt2")))
            hides_on_click = re.search(r'presetClass="exit"[^>]*>(?:(?!presetClass).)*?spid="%d"' % sh.shape_id, timing, re.S)
            boxy = geom is None or geom.get("prst", "rect") in ("rect", "roundRect", "snipRect", "flowChartProcess")
            opaque = fill is not None and not list(fill.iter(A + "alpha"))
            if not (whiteish or hides_on_click or (boxy and opaque)):
                kind = "MARK"  # coloured shape drawn on the picture (triangle, ellipse, label)
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame and len(sh.text_frame.text.strip()) > 2:
            continue
        b = bbox(sh)
        if not b:
            continue
        area = max(1e-9, (b[2] - b[0]) * (b[3] - b[1]))
        if area < 0.002 and kind == "COVER":  # tiny filled dots are annotations, not covers
            kind = "MARK"
        for p in pics:
            pb = p["bbox"]
            ix = max(0, min(b[2], pb[2]) - max(b[0], pb[0])) * max(0, min(b[3], pb[3]) - max(b[1], pb[1]))
            if ix / area > 0.5:
                anim = ""
                m = re.search(r'presetClass="(exit|entr)"[^>]*>(?:(?!presetClass).)*?spid="%d"' % sh.shape_id, timing, re.S)
                if m:
                    anim = " animation: " + ("disappears on click (hides something until then)" if m.group(1) == "exit" else "appears on click")
                tgt = p["file"] if p["file"].startswith("(") else f"IMAGE media/{p['file']}"
                what = "" if kind == "COVER" else " (annotation: reproduce it, e.g. crop the region from the reference page)"
                out_lines.append(f"- {kind} '{sh.name}' over {tgt} {pos(sh)}{anim}{what}")
                break
    return out_lines


def shape_label(slide, spid):
    for sh in slide.shapes:
        stack = [sh]
        while stack:
            x = stack.pop()
            if x.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack.extend(x.shapes)
            if x.shape_id == spid:
                t = x.text_frame.text.strip().replace("\n", " / ")[:50] if getattr(x, "has_text_frame", False) and x.has_text_frame else ""
                kind = "picture" if hasattr(x, "image") else ("group" if x.shape_type == MSO_SHAPE_TYPE.GROUP else "shape")
                return f"{kind} '{x.name}'" + (f" \"{t}\"" if t else "")
    return f"shape id {spid}"


def click_steps(slide):
    """Click steps from p:timing: for each click, which shapes appear (entr), disappear
    (exit) or are emphasised. The deck must reproduce these with v-click / v-after /
    v-switch; otherwise answers that PowerPoint shows only on a click are visible at once."""
    P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    timing = slide._element.find(P_NS + "timing")
    if timing is None:
        return []
    steps, cur = [], None
    for ctn in timing.iter(P_NS + "cTn"):
        nt = ctn.get("nodeType")
        if nt not in ("clickEffect", "withEffect", "afterEffect"):
            continue
        cls = ctn.get("presetClass", "")
        ids = sorted({int(t.get("spid")) for t in ctn.iter(P_NS + "spTgt") if t.get("spid", "").isdigit()})
        if not ids:
            continue
        if nt == "clickEffect" or cur is None:
            cur = []
            steps.append(cur)
        verb = {"entr": "show", "exit": "hide", "emph": "emphasise", "path": "move"}.get(cls, cls or "animate")
        for i in ids:
            cur.append(f"{verb} {shape_label(slide, i)}" + ("" if nt != "afterEffect" else " (after previous)"))
    return steps


def title_of(slide):
    t = slide.shapes.title
    if t is not None and t.has_text_frame and t.text_frame.text.strip():
        return t.text_frame.text.strip().replace("\n", " ")
    return ""


index = []
pictures = []
vis = 0
masters = [m.part.partname for m in prs.slide_masters]
for i, slide in enumerate(prs.slides, 1):
    if slide._element.get("show") == "0":
        continue
    vis += 1
    title = title_of(slide)
    lines.append(f"\n## Slide {vis} (original #{i}) — title: {title or '(none)'}")
    mi = masters.index(slide.slide_layout.slide_master.part.partname)
    lines.append(f"layout: {slide.slide_layout.name} (master {mi})")
    for sh in slide.shapes:
        dump_shape(sh, vis)
    lines.extend(covers(slide, vis))
    raw = raw_text(slide._element)
    lines.append("raw text: " + raw)
    rels = [r.target_ref for r in slide.part.rels.values() if r.is_external]
    if rels:
        lines.append("links: " + " ; ".join(rels))
    steps = click_steps(slide)
    for k, st in enumerate(steps, 1):
        lines.append(f"CLICK {k}: " + "; ".join(dict.fromkeys(st)))
    if steps:
        lines.append(f"clicks: {len(steps)} -> reproduce with v-click/v-after/v-switch (answers must stay hidden until the click)")
    notes = ""
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            lines.append("notes: " + notes.replace("\n", " / "))
    index.append({"slide": vis, "original": i, "title": title, "text": raw, "notes": notes, "clicks": len(steps),
                  "source": os.path.abspath(src), "source_mtime": os.path.getmtime(src)})

with open(f"{out}/dump.md", "w") as f:
    f.write(f"# {os.path.basename(src)} — {vis} visible slides of {len(prs.slides)}, "
            f"size {W}x{H} EMU (ratio {W/H:.3f})\n" + "\n".join(lines) + "\n")
with open(f"{out}/slides.json", "w") as f:
    json.dump(index, f, ensure_ascii=False, indent=1)
with open(f"{out}/pictures.json", "w") as f:
    json.dump(pictures, f, ensure_ascii=False, indent=1)
print(f"{vis} visible slides ({len(prs.slides) - vis} hidden) -> {out}/dump.md, slides.json, media/")
